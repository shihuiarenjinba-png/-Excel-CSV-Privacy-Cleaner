import csv
import io
from pathlib import Path
import tempfile
from datetime import datetime, timezone
from openpyxl import Workbook, load_workbook
from openpyxl.comments import Comment

from windows_app import clean_csv, clean_xlsx

with tempfile.TemporaryDirectory() as tmp:
    root = Path(tmp)
    csv_src = root / "sample.csv"
    csv_src.write_text("name,email\nAlice,alice@example.com\n", encoding="utf-8")
    csv_out = root / "safe.csv"
    clean_csv(csv_src, csv_out, {"email"}, "mask", "salt")
    text = csv_out.read_text(encoding="utf-8-sig")
    assert "alice@example.com" not in text

    # Regression: quoted embedded newlines, commas and quotes in unselected
    # columns must survive a read/transform/write round trip as cell values.
    mixed_headers = ["email", "keep_text", "keep_emoji", "keep_zero", "keep_empty", "keep_space"]
    mixed_row = {
        "email": "holdout@example.com",
        "keep_text": 'English, 日本語, "quoted"\nsecond line',
        "keep_emoji": "沖縄🌺🙂",
        "keep_zero": "001234",
        "keep_empty": "",
        "keep_space": "  preserve surrounding whitespace  ",
    }
    for label, encoding in (("utf8", "utf-8"), ("utf8-bom", "utf-8-sig"), ("utf16", "utf-16")):
        source = root / f"mixed-{label}.csv"
        buffer = io.StringIO(newline="")
        writer = csv.DictWriter(buffer, fieldnames=mixed_headers, lineterminator="\n")
        writer.writeheader()
        writer.writerow(mixed_row)
        source.write_bytes(buffer.getvalue().encode(encoding))
        output = root / f"mixed-{label}-safe.csv"
        clean_csv(source, output, {"email"}, "mask", "salt")
        with output.open("r", encoding="utf-8-sig", newline="") as fh:
            cleaned = next(csv.DictReader(fh))
        assert cleaned["email"] == "[REDACTED_EMAIL]"
        for column in mixed_headers[1:]:
            assert cleaned[column] == mixed_row[column], (label, column, cleaned[column], mixed_row[column])

    cp932_row = dict(mixed_row)
    cp932_row["keep_emoji"] = "沖縄"
    cp932_source = root / "mixed-cp932.csv"
    cp932_buffer = io.StringIO(newline="")
    cp932_writer = csv.DictWriter(cp932_buffer, fieldnames=mixed_headers, lineterminator="\n")
    cp932_writer.writeheader()
    cp932_writer.writerow(cp932_row)
    cp932_source.write_bytes(cp932_buffer.getvalue().encode("cp932"))
    cp932_output = root / "mixed-cp932-safe.csv"
    clean_csv(cp932_source, cp932_output, {"email"}, "mask", "salt")
    with cp932_output.open("r", encoding="utf-8-sig", newline="") as fh:
        cp932_cleaned = next(csv.DictReader(fh))
    assert cp932_cleaned["email"] == "[REDACTED_EMAIL]"
    for column in mixed_headers[1:]:
        assert cp932_cleaned[column] == cp932_row[column], ("cp932", column)

    bad_csv = root / "undecodable.csv"
    bad_csv.write_bytes(b"\x81")
    try:
        clean_csv(bad_csv, root / "bad.csv", set(), "mask", "salt")
    except UnicodeError:
        pass
    else:
        raise AssertionError("Undecodable CSV must be rejected")

    binary_csv = root / "binary.csv"
    binary_csv.write_bytes(b"\x00" * 64)
    try:
        clean_csv(binary_csv, root / "binary-out.csv", set(), "mask", "salt")
    except UnicodeError:
        pass
    else:
        raise AssertionError("Binary-like CSV must be rejected")

    duplicate_csv = root / "duplicate.csv"
    duplicate_csv.write_text("name,email,email\nAlice,a@example.com,b@example.com\n", encoding="utf-8")
    try:
        clean_csv(duplicate_csv, root / "duplicate-out.csv", set(), "mask", "salt")
    except ValueError as exc:
        assert "Duplicate" in str(exc)
    else:
        raise AssertionError("Duplicate CSV headers must be rejected")

    ragged_csv = root / "ragged.csv"
    ragged_csv.write_text("name,email\nAlice,a@example.com,EXTRA\n", encoding="utf-8")
    try:
        clean_csv(ragged_csv, root / "ragged-out.csv", set(), "mask", "salt")
    except ValueError as exc:
        assert "more fields" in str(exc)
    else:
        raise AssertionError("CSV rows wider than the header must be rejected")

    overwrite_csv = root / "overwrite.csv"
    overwrite_csv.write_text("name,email\nAlice,a@example.com\n", encoding="utf-8")
    try:
        clean_csv(overwrite_csv, overwrite_csv, {"email"}, "mask", "salt")
    except ValueError as exc:
        assert "overwrite" in str(exc).lower()
    else:
        raise AssertionError("Original CSV overwrite must be rejected")

    xlsx_src = root / "sample.xlsx"
    wb = Workbook()
    wb.properties.creator = "Private Author"
    wb.properties.lastModifiedBy = "Private Editor"
    wb.properties.created = datetime(2001, 2, 3, tzinfo=timezone.utc)
    ws = wb.active
    ws.append(["name", "email", "external_formula", "external_link"])
    ws.append(["Alice", "alice@example.com", "='[secret.xlsx]Sheet1'!A1", "portal"])
    ws["B2"].comment = Comment("private note", "Private Author")
    ws["D2"].hyperlink = "https://example.invalid/private/path"
    hidden = wb.create_sheet("Hidden")
    hidden.sheet_state = "hidden"
    wb.save(xlsx_src)

    xlsx_out = root / "safe.xlsx"
    stats = clean_xlsx(xlsx_src, xlsx_out, {"email"}, "pseudonymize", "salt", True, True)
    out_wb = load_workbook(xlsx_out, keep_links=False)
    out_ws = out_wb.active

    assert "Hidden" not in out_wb.sheetnames
    assert out_ws["B2"].value != "alice@example.com"
    assert out_ws["B2"].comment is None
    assert out_ws["C2"].value == "[REMOVED_EXTERNAL_FORMULA]"
    assert out_ws["D2"].hyperlink is None
    assert out_wb.properties.creator == "openpyxl"
    assert out_wb.properties.lastModifiedBy in (None, "")
    assert out_wb.properties.created == datetime(1980, 1, 1)
    assert stats["removed_external_formulas"] == 1
    assert stats["removed_external_hyperlinks"] == 1
    assert not getattr(out_wb, "_external_links", [])

    duplicate_xlsx = root / "duplicate.xlsx"
    dup_wb = Workbook()
    dup_ws = dup_wb.active
    dup_ws.append(["name", "email", "email"])
    dup_ws.append(["Alice", "a@example.com", "b@example.com"])
    dup_wb.save(duplicate_xlsx)
    try:
        clean_xlsx(duplicate_xlsx, root / "duplicate-safe.xlsx", set(), "mask", "salt", True, True)
    except ValueError as exc:
        assert "Duplicate" in str(exc)
    else:
        raise AssertionError("Duplicate XLSX headers must be rejected")

print("Excel CSV Privacy Cleaner smoke test: OK")
