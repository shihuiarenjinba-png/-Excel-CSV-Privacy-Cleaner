"""Selected formula values must never become live partially masked formulas."""
import tempfile
import unittest
from pathlib import Path

from openpyxl import Workbook, load_workbook
from windows_app import clean_xlsx


class ExcelFormulaExportTests(unittest.TestCase):
    def test_partial_mask_on_formula_is_rejected_without_output(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            source = root / "input.xlsx"
            output = root / "masked.xlsx"
            wb = Workbook()
            ws = wb.active
            ws.append(["name", "notes"])
            ws.append(["=HYPERLINK(\"https://example.invalid\", \"Click\")", "keep"])
            wb.save(source)
            original = source.read_bytes()
            with self.assertRaisesRegex(ValueError, "cannot be partially masked"):
                clean_xlsx(source, output, {"name"}, "mask", "salt", True, True)
            self.assertFalse(output.exists())
            self.assertEqual(source.read_bytes(), original)

    def test_pseudonymization_does_not_preserve_formula(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            source = root / "input.xlsx"
            dest = root / "pseudonym.xlsx"
            wb = Workbook()
            ws = wb.active
            ws.append(["name", "keep"])
            ws.append(["=1+1", "ordinary"])
            wb.save(source)
            clean_xlsx(source, dest, {"name"}, "pseudonymize", "private-key", True, True)
            check = load_workbook(dest, data_only=False)
            self.assertEqual(check.active["A2"].data_type, "s")
            self.assertTrue(check.active["A2"].value.startswith("USER-"))
            self.assertEqual(check.active["B2"].value, "ordinary")
            check.close()


if __name__ == "__main__":
    unittest.main()
