"""Commercial fail-closed checks: workbook structure and formulas on XLSX drops."""
import tempfile
import unittest
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.worksheet.table import Table, TableStyleInfo

from windows_app import clean_xlsx


class XlsxComplexWorkbookTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.dir = Path(self.temp.name)
        self.src = self.dir / "source.xlsx"
        self.dst = self.dir / "out.xlsx"

    def assert_rejected_without_damage(self, wb, expected_error, *, mode="mask"):
        wb.save(self.src)
        original = self.src.read_bytes()
        with self.assertRaisesRegex(ValueError, expected_error):
            clean_xlsx(self.src, self.dst, {"email"}, mode, "key", True, True)
        self.assertEqual(self.src.read_bytes(), original)
        self.assertFalse(self.dst.exists())

    def test_xlsx_drop_rejects_formula_references(self):
        wb = Workbook()
        ws = wb.active
        ws.append(["email", "amount", "calc"])
        ws.append(["a@example.com", 20, "=B2*2"])
        self.assert_rejected_without_damage(wb, "Dropping columns can break formula references", mode="drop")

    def test_xlsx_edit_rejects_merged_cells(self):
        wb = Workbook()
        ws = wb.active
        ws.append(["email", "memo"])
        ws.append(["a@example.com", "hello"])
        ws.merge_cells("A4:B4")
        self.assert_rejected_without_damage(wb, "simple data worksheets")

    def test_xlsx_edit_rejects_structured_excel_tables(self):
        wb = Workbook()
        ws = wb.active
        ws.append(["email", "memo"])
        ws.append(["a@example.com", "hello"])
        table = Table(displayName="InputTable", ref="A1:B2")
        table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
        ws.add_table(table)
        self.assert_rejected_without_damage(wb, "simple data worksheets")

    def test_drop_ignores_formulas_only_in_removed_hidden_sheets(self):
        wb = Workbook()
        ws = wb.active
        ws.append(["email", "memo"])
        ws.append(["a@example.com", "keep"])
        hidden = wb.create_sheet("Hidden")
        hidden.append(["foo"])
        hidden.append(["=SUM(1,2)"])
        hidden.sheet_state = "hidden"
        wb.save(self.src)
        result = clean_xlsx(self.src, self.dst, {"email"}, "drop", "key", True, True)
        self.assertEqual(result["removed_hidden_sheets"], 1)
        check = load_workbook(self.dst)
        try:
            self.assertEqual(check.sheetnames, [ws.title])
            self.assertEqual(check.active["A1"].value, "memo")
            self.assertEqual(check.active["A2"].value, "keep")
        finally:
            check.close()


if __name__ == "__main__":
    unittest.main()
