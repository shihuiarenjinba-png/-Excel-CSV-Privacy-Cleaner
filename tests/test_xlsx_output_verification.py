"""XLSX output must reopen correctly; never publish a broken result."""
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch
from zipfile import BadZipFile

from openpyxl import Workbook, load_workbook

from windows_app import clean_xlsx


class XlsxOutputVerificationTests(unittest.TestCase):
    def make_source(self, directory):
        src = directory / "source.xlsx"
        book = Workbook()
        sheet = book.active
        sheet.title = "People"
        sheet.append(["name", "email"])
        sheet.append(["Alice", "alice@example.com"])
        sheet2 = book.create_sheet("Other")
        sheet2.append(["count"])
        sheet2.append([4])
        book.save(src)
        return src

    def test_valid_output_reopens_and_has_same_sheets(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            src = self.make_source(root)
            out = root / "public.xlsx"
            before = src.read_bytes()
            stats = clean_xlsx(src, out, {"email"}, "mask", "salt", True, True)
            self.assertEqual(stats["sheets"], 2)
            check = load_workbook(out, read_only=True)
            try:
                self.assertEqual(check.sheetnames, ["People", "Other"])
                self.assertEqual(check["People"]["B2"].value, "[REDACTED_EMAIL]")
                self.assertEqual(check["Other"]["A2"].value, 4)
            finally:
                check.close()
            self.assertEqual(src.read_bytes(), before)

    def test_absent_selected_columns_do_not_create_misleading_xlsx(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            source = self.make_source(root)
            target = root / "should-not-exist.xlsx"
            original = source.read_bytes()
            with self.assertRaisesRegex(ValueError, "No selected columns matched"):
                clean_xlsx(source, target, {"unavailable"}, "mask", "key", True, True)
            self.assertFalse(target.exists())
            self.assertEqual(source.read_bytes(), original)

    def test_corrupt_save_is_deleted_after_verification_fails(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            src = self.make_source(root)
            out = root / "broken.xlsx"
            before = src.read_bytes()

            def malicious_or_corrupt_save(_workbook, stream):
                stream.write(b"this is not an xlsx file")

            with patch("openpyxl.workbook.workbook.Workbook.save",
                       malicious_or_corrupt_save):
                with self.assertRaises(BadZipFile):
                    clean_xlsx(src, out, {"email"}, "mask", "salt", True, True)
            self.assertFalse(out.exists())
            self.assertEqual(src.read_bytes(), before)

    def test_matching_dimensions_with_changed_saved_cell_is_rejected(self):
        from openpyxl import load_workbook as actual_loader
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            src = self.make_source(root)
            out = root / "wrong-cell.xlsx"
            before = src.read_bytes()

            def fake_reopen(path, *args, **kwargs):
                if Path(path).resolve() == out.resolve():
                    kwargs["read_only"] = False
                    opened = actual_loader(path, *args, **kwargs)
                    opened["People"]["B2"] = "alice@example.com"
                    return opened
                return actual_loader(path, *args, **kwargs)

            with patch("openpyxl.load_workbook", side_effect=fake_reopen):
                with self.assertRaisesRegex(ValueError, "row 2"):
                    clean_xlsx(src, out, {"email"}, "mask", "salt", True, True)
            self.assertFalse(out.exists())
            self.assertEqual(src.read_bytes(), before)



if __name__ == "__main__":
    unittest.main()
