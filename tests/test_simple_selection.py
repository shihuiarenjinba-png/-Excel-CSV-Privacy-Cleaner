"""Regression cases for simple checked-column selection, preview and output safety."""
import csv
import tempfile
import unittest
from pathlib import Path

from openpyxl import Workbook, load_workbook
from windows_app import available_columns, preview_changes, clean_csv, clean_xlsx


class SimpleSelectionTests(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.addCleanup(self.tmp.cleanup)
        self.root = Path(self.tmp.name)

    def test_csv_unicode_columns_and_preview_are_non_mutating(self):
        src = self.root / "顧客 CSV.csv"
        with src.open("w", encoding="utf-8-sig", newline="") as stream:
            csv.writer(stream).writerows([
                ["名前", "email", "備考"],
                ["Alice", "alice@example.com", "日本語\nquoted"],
            ])
        raw = src.read_bytes()
        self.assertEqual(available_columns(src), ["名前", "email", "備考"])
        lines = preview_changes(src, {"email"}, "mask", "local-key")
        self.assertTrue(any("[REDACTED_EMAIL]" in line for line in lines))
        self.assertEqual(src.read_bytes(), raw)
        self.assertTrue(any("metadata and formulas" in line for line in lines))

    def test_heterogeneous_sheets_union_for_checkbox_list(self):
        path = self.root / "test.xlsx"
        wb = Workbook()
        first = wb.active
        first.title = "People"
        first.append(["name", "email"])
        first.append(["Alice", "alice@example.com"])
        other = wb.create_sheet("Orders")
        other.append(["order", "amount"])
        other.append(["A1", 100])
        hidden = wb.create_sheet("Private")
        hidden.sheet_state = "hidden"
        hidden.append(["internal_secret"])
        wb.save(path)
        self.assertEqual(available_columns(path), ["name", "email", "order", "amount"])
        sample = preview_changes(path, {"email"}, "mask", "salt")
        self.assertTrue(any("[People] email:" in line for line in sample))

    def test_existing_csv_output_never_replaced(self):
        source = self.root / "source.csv"
        source.write_text("name,email\nA,a@example.com\n", encoding="utf-8")
        output = self.root / "old.csv"
        output.write_bytes(b"existing valuable work")
        with self.assertRaises(FileExistsError):
            clean_csv(source, output, {"email"}, "mask", "salt")
        self.assertEqual(output.read_bytes(), b"existing valuable work")

    def test_existing_xlsx_output_never_replaced(self):
        source = self.root / "source.xlsx"
        wb = Workbook()
        wb.active.append(["name", "email"])
        wb.active.append(["Alice", "alice@example.com"])
        wb.save(source)
        output = self.root / "existing.xlsx"
        output.write_bytes(b"existing valuable work")
        with self.assertRaises(FileExistsError):
            clean_xlsx(source, output, {"email"}, "mask", "salt", True, True)
        self.assertEqual(output.read_bytes(), b"existing valuable work")
        actual = self.root / "actual.xlsx"
        clean_xlsx(source, actual, {"email"}, "mask", "salt", True, True)
        after = load_workbook(actual, read_only=True)
        self.assertEqual(after.active["B2"].value, "[REDACTED_EMAIL]")
        after.close()

    def test_duplicate_headers_rejected_before_selection(self):
        source = self.root / "duplicate.csv"
        source.write_text("email,email\none,two\n", encoding="utf-8")
        with self.assertRaises(ValueError):
            available_columns(source)


if __name__ == "__main__":
    unittest.main()
