from __future__ import annotations

import csv
import hashlib
import hmac
import io
import re
from itertools import zip_longest
import secrets
from datetime import datetime, timezone
from zipfile import ZipFile, BadZipFile
from pathlib import Path
from tkinter import BOTH, END, LEFT, Button, Canvas, Checkbutton, Entry, Frame, IntVar, Label, Listbox, Radiobutton, Scrollbar, StringVar, Tk, filedialog, messagebox
from tkinter import ttk
from ui_export_approval import export_revision
from ui_desktop import bind_shortcuts, restore_window_state, save_window_state
from ui_commandbar import CommandBar

EMAIL_RE = re.compile(r"[A-Z0-9._%+-]+@[A-Z0-9.-]+\.[A-Z]{2,}", re.I)
PHONE_RE = re.compile(r"(?<!\d)(?:\+?81[- ]?)?0?\d{1,4}[- ]?\d{1,4}[- ]?\d{3,4}(?!\d)")
EXTERNAL_FORMULA_RE = re.compile(r"\[[^\]]+\.(?:xls|xlsx|xlsm|xlsb|csv)\]", re.I)
TEXT_ENCODINGS = ("utf-8-sig", "utf-8", "cp932", "utf-16")
MODES = {"mask", "pseudonymize", "drop"}
METADATA_EPOCH = datetime(1980, 1, 1, tzinfo=timezone.utc)
MAX_CSV_BYTES = 25 * 1024 * 1024
MAX_XLSX_BYTES = 15 * 1024 * 1024
MAX_XLSX_UNCOMPRESSED = 160 * 1024 * 1024
MAX_SPREADSHEET_ROWS = 100_000


def pseudonym(value: object, salt: str) -> str:
    """Stable keyed IDs with a 64-bit display digest.

    These are pseudonyms, not irreversible anonymization. Keep the key private.
    """
    if not salt.strip():
        raise ValueError("Pseudonymization requires a nonempty private key.")
    key = salt.encode("utf-8", errors="strict")
    raw = str(value).encode("utf-8", errors="strict")
    return "USER-" + hmac.new(key, raw, hashlib.sha256).hexdigest()[:16].upper()


def mask_value(value: object) -> str:
    text = "" if value is None else str(value)
    if EMAIL_RE.fullmatch(text):
        return "[REDACTED_EMAIL]"
    if PHONE_RE.fullmatch(text):
        return "[REDACTED_PHONE]"
    if len(text) <= 2:
        return "*" * len(text)
    return text[:1] + "*" * max(1, len(text) - 2) + text[-1:]


def parse_columns(text: str) -> set[str]:
    return {c.strip() for c in text.split(",") if c.strip()}


def validate_decoded_text(text: str) -> str:
    if "\x00" in text:
        raise UnicodeError("Decoded text contains NUL bytes and appears binary.")
    if text:
        controls = sum(ord(ch) < 32 and ch not in "\t\n\r\f" for ch in text)
        if controls / len(text) > 0.01:
            raise UnicodeError("Decoded text contains too many control characters and appears binary.")
    return text


def decode_text(raw: bytes) -> str:
    last_error: UnicodeError | None = None
    for enc in TEXT_ENCODINGS:
        try:
            return validate_decoded_text(raw.decode(enc))
        except UnicodeError as exc:
            last_error = exc
    raise UnicodeError(f"Unsupported or undecodable CSV encoding; supported encodings: {', '.join(TEXT_ENCODINGS)}") from last_error


def ensure_distinct_paths(src: Path, out: Path) -> None:
    try:
        same = src.resolve() == out.resolve()
    except OSError:
        same = src.absolute() == out.absolute()
    if same:
        raise ValueError("Refusing to overwrite the original file.")


def validate_headers(raw_headers: list[object], *, label: str) -> list[str]:
    headers = ["" if value is None else str(value).strip() for value in raw_headers]
    if not headers or not any(headers):
        raise ValueError(f"{label} header row is missing or could not be parsed.")
    if any(not header for header in headers):
        raise ValueError(f"{label} contains a blank header. Every processed column must have a unique name.")
    seen: set[str] = set()
    duplicates: list[str] = []
    for header in headers:
        if header in seen and header not in duplicates:
            duplicates.append(header)
        seen.add(header)
    if duplicates:
        raise ValueError(f"Duplicate {label} header(s): " + ", ".join(duplicates))
    return headers


def validate_input_size(path: Path) -> None:
    """Reject unsupported sizes and suspicious ZIP containers before any GUI/export read."""
    ext = path.suffix.lower()
    if ext not in {".csv", ".xlsx"}:
        raise ValueError("Only CSV and XLSX are supported.")
    if path.is_symlink() or not path.is_file():
        raise ValueError("A regular, non-symlink input file is required.")
    size = path.stat().st_size
    cap = MAX_CSV_BYTES if ext == ".csv" else MAX_XLSX_BYTES
    if size <= 0 or size > cap:
        raise ValueError(f"Input file must be between 1 byte and {cap // (1024 * 1024)} MiB.")
    if ext == ".xlsx":
        try:
            with ZipFile(path) as archive:
                members = archive.infolist()
                if len(members) > 10_000:
                    raise ValueError("Workbook contains too many ZIP entries.")
                if sum(item.file_size for item in members) > MAX_XLSX_UNCOMPRESSED:
                    raise ValueError("Expanded workbook is too large for safe local processing.")
                if any(item.filename.lower().endswith("vbaproject.bin") for item in members):
                    raise ValueError("Macro-containing spreadsheet is not supported.")
        except BadZipFile as exc:
            raise ValueError("XLSX container is damaged.") from exc


def csv_formula_risk(value: object) -> bool:
    """Conservatively block spreadsheet-formula-like CSV cells before export.

    CSV can become executable formulas when opened by spreadsheet software.
    This checks the actual output values; it does not mutate unselected cells.
    """
    if value is None:
        return False
    content = str(value).lstrip(" \t\r\n")
    if not content:
        return False
    if content[0] in ("=", "+", "@"):
        return True
    if content.startswith("-"):
        return not bool(re.fullmatch(r"-[0-9]+(?:[.][0-9]+)?", content))
    return False


def clean_csv(src: Path, out: Path, columns: set[str], mode: str, salt: str) -> dict:
    """Preflight every CSV cell, then write without keeping row dictionaries in RAM.

    The second pass uses the exact same decoded text, so the output cannot differ
    from the inspected source due to a concurrent change on disk.
    """
    ensure_distinct_paths(src, out)
    if mode not in MODES:
        raise ValueError("Mode must be mask, pseudonymize, or drop.")
    if mode == "pseudonymize" and not salt.strip():
        raise ValueError("Pseudonymization requires a nonempty private key.")
    validate_input_size(src)
    decoded = decode_text(src.read_bytes())

    def new_reader():
        reader = csv.DictReader(io.StringIO(decoded, newline=""))
        # Header normalization must apply to the actual row keys as well as
        # the checkbox labels. Otherwise " email " silently becomes blank.
        reader.fieldnames = validate_headers(list(reader.fieldnames or []), label="CSV")
        return reader

    reader = new_reader()
    headers = list(reader.fieldnames or [])
    if any(csv_formula_risk(header) for header in headers):
        raise ValueError("CSV column header could execute as a spreadsheet formula. No file exported.")
    selected = columns or set(headers)
    unknown = sorted(selected - set(headers))
    if columns and unknown:
        raise ValueError("Requested column(s) not found: " + ", ".join(unknown))
    if mode == "drop":
        out_headers = [h for h in headers if h not in selected]
        if not out_headers:
            raise ValueError("Refusing to drop every CSV column.")
    else:
        out_headers = headers

    def transformed():
        for row_index, row in enumerate(new_reader(), start=2):
            if row_index - 1 > MAX_SPREADSHEET_ROWS:
                raise ValueError("CSV exceeds 100,000 rows; split the file before processing.")
            if None in row:
                raise ValueError(f"CSV row has more fields than the header at line {row_index}.")
            out_row = dict(row)
            for col in selected:
                if mode == "drop":
                    out_row.pop(col, None)
                elif mode == "pseudonymize":
                    original_value = out_row.get(col)
                    # Missing/blank values are not people. Pseudonymizing them
                    # would invent nonempty IDs and corrupt downstream counts.
                    out_row[col] = (
                        "" if original_value in (None, "")
                        else pseudonym(original_value, salt)
                    )
                else:
                    out_row[col] = mask_value(out_row.get(col, ""))
            for column in out_headers:
                if csv_formula_risk(out_row.get(column)):
                    raise ValueError(
                        f"Possible spreadsheet formula in output row {row_index}, column '{column}'. "
                        "Select and pseudonymize/drop that column or remove the risky value. "
                        "No file was exported."
                    )
            yield {k: (out_row.get(k) if out_row.get(k) is not None else "") for k in out_headers}

    # Check all rows BEFORE opening a new output. Fail closed on late-row errors.
    count = 0
    for _row in transformed():
        count += 1

    created = False
    try:
        with out.open("x", encoding="utf-8-sig", newline="") as fh:
            created = True
            writer = csv.DictWriter(fh, fieldnames=out_headers)
            writer.writeheader()
            for out_row in transformed():
                writer.writerow(out_row)
        # Independently parse the published bytes. A writer or storage failure
        # must never leave a plausible but silently corrupted "safe" copy.
        with out.open("r", encoding="utf-8-sig", newline="") as verify_handle:
            saved = csv.DictReader(verify_handle)
            if saved.fieldnames != out_headers:
                raise ValueError("CSV output verification failed: unexpected header.")
            for index, (saved_row, expected_row) in enumerate(
                zip_longest(saved, transformed()), start=2
            ):
                if saved_row is None or expected_row is None or saved_row != expected_row:
                    raise ValueError(f"CSV output verification failed at row {index}.")
    except BaseException:
        if created:
            out.unlink(missing_ok=True)
        raise
    return {"rows": count, "columns": len(headers)}


def is_external_formula(value: object) -> bool:
    if not isinstance(value, str) or not value.startswith("="):
        return False
    lower = value.lower()
    return bool(EXTERNAL_FORMULA_RE.search(value) or "http://" in lower or "https://" in lower or "\\\\" in value)


def scrub_core_properties(wb) -> None:
    props = wb.properties
    # openpyxl requires valid datetimes when serializing core created/modified
    # properties. A fixed UTC epoch removes source timestamps without making the
    # generated XLSX unsaveable on current supported openpyxl versions. Current
    # openpyxl rewrites `modified` to the save time, which also prevents the
    # source document's value from persisting.
    for attr in ("created", "modified"):
        if hasattr(props, attr):
            setattr(props, attr, METADATA_EPOCH)
    for attr in (
        "creator", "lastModifiedBy", "title", "subject", "description", "keywords",
        "category", "contentStatus", "identifier", "language", "version", "revision",
        "lastPrinted"
    ):
        if hasattr(props, attr):
            setattr(props, attr, None)



def validate_xlsx_edit_scope(wb, columns: set[str], mode: str, remove_hidden: bool) -> None:
    """Refuse workbook features that openpyxl cannot reliably remap on column edits.

    Most spreadsheet structures are *not* equivalent to a plain data table.
    In particular, delete_cols does not rewrite formula references.
    """
    edited_columns: dict[str, set[int]] = {}
    for ws in wb.worksheets:
        if remove_hidden and ws.sheet_state != "visible":
            continue
        matches = {
            cell.column for cell in ws[1]
            if cell.value is not None and str(cell.value).strip() in columns
        }
        if matches:
            edited_columns[ws.title] = matches
    for ws in wb.worksheets:
        if remove_hidden and ws.sheet_state != "visible":
            continue
        has_drawings = bool(ws._charts or ws._images or ws._pivots)
        has_structures = bool(ws.tables or ws.merged_cells.ranges)
        if has_drawings or has_structures:
            raise ValueError(
                f"Worksheet {ws.title!r} contains tables, merged cells, "
                "images, charts or pivot tables. This version only supports "
                "simple data worksheets. No file was exported."
            )
    if mode in {"mask", "pseudonymize"} and edited_columns:
        for ws in wb.worksheets:
            if remove_hidden and ws.sheet_state != "visible":
                continue
            for row in ws.iter_rows():
                for cell in row:
                    if cell.data_type != "f" or is_external_formula(cell.value):
                        continue
                    if ws.title in edited_columns and cell.column in edited_columns[ws.title] and cell.row >= 2:
                        # This formula is itself selected for replacement.
                        # Pseudonymize replaces the WHOLE formula, so no broken
                        # live formula survives. Partial mask is rejected later
                        # by clean_xlsx with its specific user-facing message.
                        continue
                    if ws.title in edited_columns or "!" in str(cell.value):
                        raise ValueError(
                            f"Worksheet {ws.title!r} has formulas that may depend on edited cells. "
                            "This version cannot verify formula dependencies. "
                            "Use a formulas-free table copy or CSV. No file was exported."
                        )
    if mode == "drop":
        if wb.defined_names:
            raise ValueError(
                "Named ranges can break when columns are dropped. "
                "Remove named ranges in a copy or use CSV. No file was exported."
            )
        for ws in wb.worksheets:
            if remove_hidden and ws.sheet_state != "visible":
                continue
            for row in ws.iter_rows():
                for cell in row:
                    if cell.data_type == "f":
                        raise ValueError(
                            f"Worksheet {ws.title!r} contains formulas. "
                            "Dropping columns can break formula references. "
                            "Remove formulas in a separate copy or use a CSV input. "
                            "No file was exported."
                        )


def clean_xlsx(src: Path, out: Path, columns: set[str], mode: str, salt: str, remove_hidden: bool, remove_comments: bool) -> dict:
    from openpyxl import load_workbook

    ensure_distinct_paths(src, out)
    if mode not in MODES:
        raise ValueError("Mode must be mask, pseudonymize, or drop.")
    if mode == "pseudonymize" and not salt.strip():
        raise ValueError("Pseudonymization requires a nonempty private key.")

    validate_input_size(src)
    wb = load_workbook(src, keep_links=False)
    validate_xlsx_edit_scope(wb, columns, mode, remove_hidden)
    scrub_core_properties(wb)
    removed_sheets = 0
    removed_external_formulas = 0
    removed_external_hyperlinks = 0
    sheet_results: list[dict] = []

    for ws in list(wb.worksheets):
        if ws.max_row > MAX_SPREADSHEET_ROWS + 1:
            raise ValueError(f"Worksheet {ws.title!r} exceeds 100,000 data rows.")
        if remove_hidden and ws.sheet_state != "visible":
            wb.remove(ws)
            removed_sheets += 1
            continue

        raw_headers = [cell.value for cell in ws[1]]
        if any(value not in (None, "") for value in raw_headers):
            header_names = validate_headers(raw_headers, label=f"worksheet '{ws.title}'")
            headers = {name: index + 1 for index, name in enumerate(header_names)}
        else:
            headers = {}

        requested = columns or set(headers.keys())
        selected = requested & set(headers.keys())
        missing = sorted(requested - set(headers.keys())) if columns else []
        sheet_result = {
            "sheet": ws.title,
            "status": "processed" if selected else "skipped_no_matching_columns",
            "processed_columns": sorted(selected),
            "missing_requested_columns": missing,
            "rows": max(0, ws.max_row - 1),
        }

        if mode == "drop":
            indices = sorted([headers[c] for c in selected if c in headers], reverse=True)
            if indices and len(indices) >= ws.max_column:
                raise ValueError(f"Refusing to drop every column from sheet '{ws.title}'.")
            for idx in indices:
                ws.delete_cols(idx, 1)
        else:
            for col in selected:
                idx = headers.get(col)
                if not idx:
                    continue
                for row in range(2, ws.max_row + 1):
                    cell = ws.cell(row=row, column=idx)
                    if cell.value in (None, ""):
                        continue
                    if mode == "mask" and isinstance(cell.value, str) and cell.value.startswith("="):
                        raise ValueError(
                            f"Selected formula cell {ws.title}!{cell.coordinate} cannot be partially masked. "
                            "Use pseudonymize/drop or remove the formula before exporting."
                        )
                    cell.value = pseudonym(cell.value, salt) if mode == "pseudonymize" else mask_value(cell.value)

        for row in ws.iter_rows():
            for cell in row:
                if remove_comments and cell.comment is not None:
                    cell.comment = None
                if is_external_formula(cell.value):
                    cell.value = "[REMOVED_EXTERNAL_FORMULA]"
                    removed_external_formulas += 1
                if cell.hyperlink is not None:
                    target = getattr(cell.hyperlink, "target", None)
                    if target and not str(target).startswith("#"):
                        cell.hyperlink = None
                        removed_external_hyperlinks += 1
        sheet_results.append(sheet_result)

    if not wb.worksheets:
        raise ValueError("All worksheets would be removed. Keep at least one visible sheet.")
    if not any(item["processed_columns"] for item in sheet_results):
        raise ValueError(
            "No selected columns matched the retained worksheets. "
            "No misleading unchanged copy was created."
        )
    created = False
    try:
        with out.open("xb") as handle:
            created = True
            wb.save(handle)
        # Saving can produce an unreadable or incomplete workbook. Refuse to
        # publish unless openpyxl can reopen it with the intended sheet layout.
        reopened = load_workbook(out, read_only=True, keep_links=False)
        try:
            if reopened.sheetnames != wb.sheetnames:
                raise ValueError("XLSX output verification failed: sheet list changed.")
            for intended in wb.worksheets:
                actual = reopened[intended.title]
                if (actual.max_row, actual.max_column) != (intended.max_row, intended.max_column):
                    raise ValueError(
                        f"XLSX output verification failed: sheet dimensions changed for {intended.title!r}."
                    )
                # Check ALL saved cell values/formulas, not only the sheet size.
                # A file can reopen with the expected dimensions while retaining
                # sensitive values or losing edits in individual cells.
                for row_number, pair in enumerate(zip_longest(
                    intended.iter_rows(values_only=True),
                    actual.iter_rows(values_only=True),
                ), start=1):
                    expected_row, saved_row = pair
                    if expected_row != saved_row:
                        raise ValueError(
                            f"XLSX output verification failed at {intended.title!r}, "
                            f"row {row_number}; no exported workbook retained."
                        )
        finally:
            reopened.close()
    except BaseException:
        if created:
            out.unlink(missing_ok=True)
        raise
    return {
        "sheets": len(wb.worksheets),
        "removed_hidden_sheets": removed_sheets,
        "removed_external_formulas": removed_external_formulas,
        "removed_external_hyperlinks": removed_external_hyperlinks,
        "sheet_results": sheet_results,
    }


def available_columns(path: Path) -> list[str]:
    """Read only headers for the selection UI; never alter the input."""
    validate_input_size(path)
    ext = path.suffix.lower()
    if ext == ".csv":
        sample = decode_text(path.read_bytes())
        reader = csv.reader(io.StringIO(sample, newline=""))
        return validate_headers(next(reader, []), label="CSV")
    if ext == ".xlsx":
        from openpyxl import load_workbook
        wb = load_workbook(path, read_only=True, keep_links=False)
        try:
            names = []
            seen = set()
            for ws in wb.worksheets:
                if ws.sheet_state != "visible":
                    continue
                first = next(ws.iter_rows(min_row=1, max_row=1, values_only=True), [])
                if not any(value not in (None, "") for value in first):
                    continue
                for name in validate_headers(list(first), label=f"worksheet '{ws.title}'"):
                    if name not in seen:
                        names.append(name)
                        seen.add(name)
            if not names:
                raise ValueError("No visible worksheet has a valid header row.")
            return names
        finally:
            wb.close()
    raise ValueError("Only CSV and XLSX inputs are supported.")


def preview_changes(path: Path, columns: set[str], mode: str, salt: str, limit: int = 4,
                    *, remove_hidden: bool = True) -> list[str]:
    """Short in-memory illustrative before/after preview; not a security audit."""
    if not columns:
        return ["Select at least one column to preview."]
    validate_input_size(path)
    rows = []
    if path.suffix.lower() == ".csv":
        decoded = decode_text(path.read_bytes())
        reader = csv.DictReader(io.StringIO(decoded, newline=""))
        reader.fieldnames = validate_headers(list(reader.fieldnames or []), label="CSV")
        for row in reader:
            if None in row:
                raise ValueError("CSV data has more fields than headers.")
            rows.append(("CSV", row))
            if len(rows) >= limit:
                break
    else:
        from openpyxl import load_workbook
        # Preview must reject unsupported structures BEFORE presenting apparently
        # valid samples that could never be exported safely.
        preflight = load_workbook(path, keep_links=False, data_only=False)
        try:
            validate_xlsx_edit_scope(preflight, columns, mode, remove_hidden)
        finally:
            preflight.close()
        wb = load_workbook(path, read_only=True, keep_links=False, data_only=False)
        try:
            for ws in wb.worksheets:
                if ws.sheet_state != "visible":
                    continue
                values = ws.iter_rows(values_only=True)
                header = next(values, ())
                if not any(value not in (None, "") for value in header):
                    continue
                names = validate_headers(list(header), label=ws.title)
                if not (set(names) & columns):
                    continue
                for values_row in values:
                    rows.append((ws.title, dict(zip(names, values_row))))
                    if len(rows) >= limit:
                        break
                if len(rows) >= limit:
                    break
        finally:
            wb.close()
    lines = ["Preview only: other sheets, metadata and formulas are not shown here."]
    for sheet, row in rows:
        for col in sorted(columns & row.keys()):
            before = "" if row.get(col) is None else str(row[col])
            after = "(column removed)" if mode == "drop" else (
                pseudonym(row[col], salt) if mode == "pseudonymize" else mask_value(row[col])
            )
            lines.append(f"[{sheet}] {col}: {before[:65]!r}  →  {after[:65]!r}")
    if len(lines) == 1:
        lines.append("No matching data rows in the preview sample.")
    return lines



def format_export_summary(output: Path, stats: dict, selected: set[str]) -> list[str]:
    """Customer-readable status. Never imply that unselected sheets were sanitized."""
    lines = [f"New copy created: {output.name}",
             f"Selected columns: {len(selected)}"]
    if output.suffix.lower() == ".csv":
        lines.append(f"CSV rows processed: {stats.get('rows', 0)}")
    else:
        sheets = stats.get("sheet_results", [])
        changed = [s for s in sheets if s.get("status") == "processed"]
        skipped = [s for s in sheets if s.get("status") != "processed"]
        lines.append(f"Worksheets changed: {len(changed)}; not changed: {len(skipped)}")
        lines.append(f"Hidden worksheets removed: {stats.get('removed_hidden_sheets', 0)}")
        for sheet in sheets[:15]:
            sheet_name = str(sheet.get("sheet", ""))
            processed = ", ".join(sheet.get("processed_columns", []))
            if sheet.get("status") == "processed":
                lines.append(f"Changed sheet: {sheet_name} | selected columns: {processed}")
            else:
                lines.append(f"Unchanged sheet (no matching columns): {sheet_name}")
        if len(sheets) > 15:
            lines.append(f"More sheet statuses: {len(sheets) - 15}; review the exported workbook.")
    lines.append("Review every output file before sharing; privacy removal is not guaranteed.")
    return lines


class App:
    def __init__(self, root: Tk):
        self.root = root
        self.root.title("Excel / CSV Privacy Cleaner — Windows")
        self.root.geometry("940x680")
        self.root.minsize(820, 600)
        restore_window_state(self.root, "ExcelCSVPrivacyCleaner", min_width=820, min_height=600)
        self.file: Path | None = None
        self.mode = StringVar(value="mask")
        self.salt = StringVar(value=secrets.token_hex(16))
        self.remove_hidden = IntVar(value=1)
        self.remove_comments = IntVar(value=1)
        self.column_vars: dict[str, IntVar] = {}
        self.column_widgets: dict[str, Checkbutton] = {}
        self.column_filter = StringVar()
        self.advanced_visible = False
        self.show_key = IntVar(value=0)
        self.key_entry = None
        self.preview_revision: str | None = None

        head = Frame(root, padx=12, pady=12)
        head.pack(fill="x")
        Label(head, text="Excel / CSV Privacy Cleaner", font=("Segoe UI", 19, "bold")).pack(anchor="w")
        Label(head, text="Select file → check columns → choose method → preview → export a new copy.").pack(anchor="w")
        Label(head, text="Local processing — the original file is never overwritten.", fg="#56616f").pack(anchor="w", pady=(3, 0))
        Label(head, text="① Choose file  →  ② Select columns  →  ③ Pick a method  →  ④ Preview  →  ⑤ Save a new copy", font=("Segoe UI", 10, "bold"), wraplength=860, justify=LEFT).pack(anchor="w", pady=(4, 0))
        self.command_bar = CommandBar(root)
        self.command_bar.pack(fill="x", padx=12, pady=(0, 6))
        self.choose_button = self.command_bar.add_command("open", "＋", "Open file", self.choose, primary=True)
        self.preview_button = self.command_bar.add_command("preview", "◉", "Preview", self.preview, enabled=False)
        self.export_button = self.command_bar.add_command("save", "▣", "Save copy", self.export, enabled=False)
        self.command_bar.add_separator()
        self.command_bar.add_command("advanced", "⚙", "Advanced", self.toggle_advanced)
        self.command_bar.add_spacer()
        Label(root, text="Open a file → choose columns and a method → Preview → Save copy. The original file is never overwritten.",
              fg="#56616f", anchor="w").pack(fill="x", padx=18, pady=(0, 6))

        Label(root, text="2  Check the columns to process (none selected by default)", anchor="w").pack(fill="x", padx=12, pady=(12, 2))
        filter_row = Frame(root, padx=12)
        filter_row.pack(fill="x", pady=(0, 4))
        Label(filter_row, text="Find columns:").pack(side=LEFT)
        Entry(filter_row, textvariable=self.column_filter, width=28).pack(side=LEFT, padx=(6, 8))
        Button(filter_row, text="Select visible", command=lambda: self.set_visible_columns(1)).pack(side=LEFT)
        Button(filter_row, text="Clear visible", command=lambda: self.set_visible_columns(0)).pack(side=LEFT, padx=(6, 0))
        column_box = Frame(root, padx=12)
        column_box.pack(fill="x")
        self.canvas = Canvas(column_box, height=165, highlightthickness=1)
        scroll = Scrollbar(column_box, orient="vertical", command=self.canvas.yview)
        self.canvas.configure(yscrollcommand=scroll.set)
        self.canvas.pack(side=LEFT, fill="both", expand=True)
        scroll.pack(side="right", fill="y")
        self.check_area = Frame(self.canvas)
        self.check_window = self.canvas.create_window((0, 0), window=self.check_area, anchor="nw")
        self.canvas.bind("<Configure>", lambda e: self.canvas.itemconfigure(self.check_window, width=e.width))
        self.check_area.bind("<Configure>", lambda _e: self.canvas.configure(scrollregion=self.canvas.bbox("all")))

        options = Frame(root, padx=12, pady=9)
        options.pack(fill="x")
        Label(options, text="3  Processing method:").pack(anchor="w")
        mode_row = Frame(options)
        mode_row.pack(fill="x")
        for description, value in (
            ("Mask (hide values)", "mask"),
            ("Pseudonymize (ID)", "pseudonymize"),
            ("Drop (remove columns)", "drop"),
        ):
            Radiobutton(mode_row, text=description, variable=self.mode,
                        value=value, anchor="w").pack(side=LEFT, padx=(0, 12))
        self.advanced = Frame(root, padx=12)
        Label(self.advanced, text="Pseudonym key (keep private; use same key for consistent IDs):").pack(anchor="w")
        self.key_entry = Entry(self.advanced, textvariable=self.salt, show="•")
        self.key_entry.pack(fill="x")
        Checkbutton(self.advanced, text="Show pseudonym key", variable=self.show_key,
                    command=self.toggle_key_visibility).pack(anchor="w")
        Checkbutton(self.advanced, text="Remove hidden worksheets (XLSX)", variable=self.remove_hidden).pack(anchor="w")
        Checkbutton(self.advanced, text="Remove cell comments (XLSX)", variable=self.remove_comments).pack(anchor="w")

        Label(root, text="Preview and processing results (example rows; verify exported file yourself):").pack(anchor="w", padx=12, pady=(12, 2))
        self.listbox = Listbox(root, font=("Consolas", 10))
        self.listbox.pack(fill=BOTH, expand=True, padx=12, pady=(0, 12))
        self.listbox.insert(END, "Choose a CSV or XLSX file to begin.")
        for variable in (self.mode, self.salt, self.remove_hidden, self.remove_comments):
            variable.trace_add("write", self._invalidate_preview)
        self.column_filter.trace_add("write", lambda *_: self._render_column_filter())
        bind_shortcuts(self.root, {
            "<Control-O>": self.choose,
            "<Control-P>": self.preview,
            "<Control-S>": self.export,
        })
        self.root.protocol("WM_DELETE_WINDOW", self._close)
        self._refresh_actions()

    def _matching_columns(self):
        query = self.column_filter.get().strip().casefold()
        return [name for name in self.column_vars if not query or query in name.casefold()]

    def _render_column_filter(self):
        visible = set(self._matching_columns())
        for name, widget in self.column_widgets.items():
            if name in visible:
                widget.pack(anchor="w")
            else:
                widget.pack_forget()

    def set_visible_columns(self, value: int):
        for name in self._matching_columns():
            self.column_vars[name].set(value)

    def _invalidate_preview(self, *_args):
        self.preview_revision = None
        self._refresh_actions()

    def _refresh_actions(self):
        ready = self.file is not None and bool(self.selected_columns())
        self.preview_button.configure(state="normal" if ready else "disabled")
        self.export_button.configure(state="normal" if ready and self.preview_revision else "disabled")

    def _close(self):
        save_window_state(self.root, "ExcelCSVPrivacyCleaner")
        self.root.destroy()

    def toggle_key_visibility(self):
        self.key_entry.configure(show="" if self.show_key.get() else "•")

    def toggle_advanced(self):
        self.advanced_visible = not self.advanced_visible
        if self.advanced_visible:
            self.advanced.pack(fill="x", before=self.listbox)
        else:
            self.advanced.pack_forget()

    def selected_columns(self) -> set[str]:
        return {name for name, flag in self.column_vars.items() if flag.get()}

    def choose(self):
        p = filedialog.askopenfilename(
            parent=self.root, title="Choose CSV or XLSX",
            filetypes=[("Data files", "*.csv *.xlsx"), ("CSV", "*.csv"), ("Excel", "*.xlsx")],
        )
        if not p:
            return
        candidate = Path(p)
        try:
            names = available_columns(candidate)
            if len(names) > 250:
                raise ValueError("Over 250 distinct columns: use a smaller input or split it before processing.")
        except Exception as exc:
            # Never leave a previously chosen source active after a failed replacement.
            self.file = None
            self.column_vars = {}
            self.column_widgets = {}
            for widget in self.check_area.winfo_children():
                widget.destroy()
            self.listbox.delete(0, END)
            self.listbox.insert(END, "Input validation failed; select a supported file before exporting.")
            messagebox.showerror("Cannot read columns", str(exc))
            return
        self.file = candidate
        self.preview_revision = None
        for widget in self.check_area.winfo_children():
            widget.destroy()
        self.column_vars = {}
        self.column_widgets = {}
        self.column_filter.set("")
        for name in names:
            state = IntVar(value=0)
            self.column_vars[name] = state
            state.trace_add("write", self._invalidate_preview)
            widget = Checkbutton(self.check_area, text=name, variable=state, anchor="w")
            self.column_widgets[name] = widget
            widget.pack(anchor="w")
        self.listbox.delete(0, END)
        self.listbox.insert(END, f"Selected: {self.file.name}")
        self.listbox.insert(END, f"Loaded {len(names)} column labels. Check only those you intend to change.")
        self._refresh_actions()

    def preview(self):
        if self.file is None:
            return messagebox.showinfo("Choose file first", "Select an input file.")
        selected = self.selected_columns()
        if not selected:
            return messagebox.showwarning("No columns selected", "Check one or more column names.")
        try:
            before = export_revision(self.file, selected, self.mode.get(), self.salt.get(),
                                     bool(self.remove_hidden.get()), bool(self.remove_comments.get()))
            lines = preview_changes(self.file, selected, self.mode.get(), self.salt.get(),
                                    remove_hidden=bool(self.remove_hidden.get()))
            after = export_revision(self.file, selected, self.mode.get(), self.salt.get(),
                                    bool(self.remove_hidden.get()), bool(self.remove_comments.get()))
            if before != after:
                raise ValueError("入力がプレビュー中に更新されました。再度確認してください。")
            self.preview_revision = after
            self._refresh_actions()
            self.listbox.delete(0, END)
            for line in lines:
                self.listbox.insert(END, line)
        except Exception as exc:
            self.preview_revision = None
            self._refresh_actions()
            messagebox.showerror("Preview failed", str(exc))

    def export(self):
        if self.file is None:
            return messagebox.showinfo("Choose file first", "Select an input file.")
        selected = self.selected_columns()
        if not selected:
            return messagebox.showwarning("No columns selected", "Check one or more column names.")
        mode = self.mode.get()
        if mode not in MODES:
            return messagebox.showerror("Invalid method", "Choose mask, pseudonymize or drop.")
        if mode == "pseudonymize" and not self.salt.get().strip():
            return messagebox.showerror("Missing key", "Pseudonymization needs a nonempty private key.")
        try:
            current_revision = export_revision(self.file, selected, mode, self.salt.get(),
                                               bool(self.remove_hidden.get()), bool(self.remove_comments.get()))
        except (ValueError, OSError) as exc:
            self.preview_revision = None
            return messagebox.showerror("Input changed", str(exc))
        if self.preview_revision != current_revision:
            return messagebox.showwarning("Preview required",
                "The input, selected columns or options have changed. Click Preview before exporting.")
        if not messagebox.askokcancel(
            "Review export",
            f"{len(selected)} selected column(s): {', '.join(sorted(selected)[:8])}"
            "\nOther columns and workbook content may still contain personal data."
            "\nWorkbook formulas and hidden structures are not comprehensively certified."
            "\nContinue to choose a NEW output file?",
        ):
            return
        ext = self.file.suffix.lower()
        out = filedialog.asksaveasfilename(
            parent=self.root, defaultextension=ext,
            initialfile=f"{self.file.stem}_safe{ext}", filetypes=[("Output", f"*{ext}")],
        )
        if not out:
            return
        destination = Path(out)
        if destination.exists() or destination.is_symlink():
            return messagebox.showerror("File exists", "Choose a new output filename; existing files are never replaced.")
        self.preview_revision = None  # Approval is single-use; no automatic replay.
        self._refresh_actions()
        try:
            if ext == ".csv":
                stats = clean_csv(self.file, destination, selected, mode, self.salt.get())
            else:
                stats = clean_xlsx(self.file, destination, selected, mode, self.salt.get(),
                                   bool(self.remove_hidden.get()), bool(self.remove_comments.get()))
            if destination.stat().st_size <= 0:
                raise ValueError("Generated file is empty; do not share.")
            if export_revision(self.file, selected, mode, self.salt.get(),
                               bool(self.remove_hidden.get()), bool(self.remove_comments.get())) != current_revision:
                # clean_csv/clean_xlsx own this brand-new output. A changed input
                # invalidates it; the customer's original source is never deleted.
                destination.unlink(missing_ok=True)
                raise ValueError("The source changed during export; discarded the new copy. Preview again.")
            self.listbox.delete(0, END)
            for line in format_export_summary(destination, stats, selected):
                self.listbox.insert(END, line)
            messagebox.showinfo("Export complete", "New copy created. Check the final output before sharing.")
        except Exception as exc:
            messagebox.showerror("Export failed", str(exc))


def main():
    from i18n import install_bilingual_ui

    root = Tk()
    App(root)
    install_bilingual_ui(root)
    root.mainloop()


if __name__ == "__main__":
    main()
